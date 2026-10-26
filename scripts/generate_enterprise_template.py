"""產生企業內部文件清單的 Excel 範本，給品保／文管單位填寫。

    python scripts/generate_enterprise_template.py

輸出到 data/templates/企業內部文件清單範本.xlsx。清單是之後「企業內規」資料來源的輸入：
系統只需要目錄（編號、名稱、版次、狀態），不需要文件全文。
"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[1] / "data" / "templates" / "企業內部文件清單範本.xlsx"
#: Arial 沒有繁體中文字形，改用微軟正黑體
FONT = "Microsoft JhengHei"

NAVY = "1F3864"
REQUIRED_FILL = PatternFill("solid", fgColor="FFF2CC")   # 必填欄底色
EXAMPLE_FILL = PatternFill("solid", fgColor="EDEDED")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def font(**kw):
    return Font(name=FONT, size=kw.pop("size", 11), **kw)


TYPES = ["品質手冊", "程序書", "作業標準（SOP）", "施工規範", "檢驗標準", "表單", "業主特訂條款", "其他"]
STATUSES = ["現行", "已修訂", "作廢"]

# (標題, 必填, 欄寬, 說明, 範例值)
COLUMNS = [
    ("文件編號", True, 14, "公司文件管制的唯一編號，同一份文件改版時不變。", "QP-07-03"),
    ("文件名稱", True, 30, "完整名稱，和工程文件裡引用時的寫法越接近越好。", "混凝土澆置作業標準"),
    ("類型", True, 16, "從下拉選單選。", "作業標準（SOP）"),
    ("版次", True, 10, "目前生效的版次，寫法照公司慣例。", "第5版"),
    ("生效日", True, 13, "目前這一版開始生效的日期。", date(2025, 3, 1)),
    ("最後修訂日", False, 13, "最近一次改版的日期；沒改過版就留空。", date(2025, 2, 20)),
    ("狀態", True, 10, "現行＝有效；已修訂＝仍有效但版次有變動；作廢＝不可再引用。", "現行"),
    ("取代文件編號", False, 14, "狀態為「作廢」時，取代它的新文件編號。", ""),
    ("主管單位", False, 14, "負責維護這份文件的部門。", "品保部"),
    ("別名", False, 32, "簡稱、舊名、常見寫法，用「、」分隔。文件編號不必重複寫。", "混凝土澆置標準、澆置SOP"),
    ("備註", False, 24, "其他說明。", "範例列，填寫時請刪除"),
]

wb = Workbook()

# --- 說明 ---
guide = wb.active
guide.title = "說明"
guide.sheet_view.showGridLines = False
guide.column_dimensions["A"].width = 16
guide.column_dimensions["B"].width = 90

rows = [
    ("企業內部文件清單（範本）", None),
    (None, None),
    ("用途", "提供給「工程法規符合性檢查平台」比對工程文件裡引用的內部文件（品質手冊、程序書、作業標準、公司版規範、業主特訂條款等）。"),
    ("要填什麼", "只需要清單，不需要文件全文。一份文件一列，填在「文件清單」工作表。"),
    ("怎麼填", "淺黃色標題的欄位必填，其餘選填。「類型」與「狀態」請用下拉選單。日期請填成日期格式（例：2025/3/1）。"),
    ("範例列", "「文件清單」第 2 列（灰底）是格式範例，填寫時請整列刪除。"),
    ("已有文件管制一覽表", "若公司已有 ISO 文件管制一覽表，可直接提供原檔，不必重抄；系統會依欄位名稱對應。請一併告知各欄位對應到本範本的哪一欄。"),
    ("誰來維護", "建議由品保／文管單位維護。每次改版或作廢都要更新，否則系統會把作廢的文件判為「現行有效」。"),
    ("機密性", "本清單只含文件名稱與版次等目錄資訊，但仍屬公司資料，提供前請依公司規定取得同意。"),
    (None, None),
    ("欄位說明", None),
]
for r, (label, text) in enumerate(rows, start=1):
    a, b = guide.cell(r, 1, label), guide.cell(r, 2, text)
    a.font = font(bold=True, color=NAVY)
    b.font = font()
    b.alignment = Alignment(wrap_text=True, vertical="top")
    a.alignment = Alignment(vertical="top")
guide["A1"].font = font(bold=True, size=16, color=NAVY)
guide["A11"].font = font(bold=True, size=13, color=NAVY)

r = 12
for title, required, _, desc, _ in COLUMNS:
    a = guide.cell(r, 1, title + ("（必填）" if required else ""))
    a.font = font(bold=required)
    if required:
        a.fill = REQUIRED_FILL
    b = guide.cell(r, 2, desc)
    b.font = font()
    b.alignment = Alignment(wrap_text=True)
    a.border = b.border = BORDER
    r += 1

r += 1
guide.cell(r, 1, "狀態對應").font = font(bold=True, size=13, color=NAVY)
r += 1
for label, meaning in [
    ("現行", "系統判為「現行有效」"),
    ("已修訂", "系統判為「已修正」：仍有效，但引用的版次可能過舊，需要人工確認"),
    ("作廢", "系統判為「已廢止」：不可再引用，報告會標示"),
]:
    guide.cell(r, 1, label).font = font()
    guide.cell(r, 2, meaning).font = font()
    guide.cell(r, 1).border = guide.cell(r, 2).border = BORDER
    r += 1

# --- 選項（下拉來源）---
opts = wb.create_sheet("選項")
opts["A1"], opts["B1"] = "類型", "狀態"
for c in ("A1", "B1"):
    opts[c].font = font(bold=True, color="FFFFFF")
    opts[c].fill = PatternFill("solid", fgColor=NAVY)
for i, t in enumerate(TYPES, start=2):
    opts.cell(i, 1, t).font = font()
for i, s in enumerate(STATUSES, start=2):
    opts.cell(i, 2, s).font = font()
opts.column_dimensions["A"].width = 18
opts.column_dimensions["B"].width = 10

# --- 文件清單 ---
ws = wb.create_sheet("文件清單", 1)
ws.freeze_panes = "A2"
LAST_ROW = 1000
for col, (title, required, width, desc, example) in enumerate(COLUMNS, start=1):
    letter = ws.cell(1, col).column_letter
    head = ws.cell(1, col, title + (" *" if required else ""))
    head.font = font(bold=True, color=NAVY if required else "FFFFFF")
    head.fill = REQUIRED_FILL if required else PatternFill("solid", fgColor=NAVY)
    head.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    head.border = BORDER
    head.comment = Comment(desc, "範本")
    ws.column_dimensions[letter].width = width

    cell = ws.cell(2, col, example if example != "" else None)
    cell.font = font(italic=True, color="595959")
    cell.fill = EXAMPLE_FILL
    cell.border = BORDER
    if isinstance(example, date):
        cell.number_format = "yyyy/mm/dd"
    for row in range(3, LAST_ROW + 1):
        c = ws.cell(row, col)
        c.font = font()
        if title in ("生效日", "最後修訂日"):
            c.number_format = "yyyy/mm/dd"
ws.row_dimensions[1].height = 30
ws.auto_filter.ref = f"A1:{ws.cell(1, len(COLUMNS)).column_letter}{LAST_ROW}"

type_dv = DataValidation(
    type="list", formula1=f"=選項!$A$2:$A${len(TYPES) + 1}", allow_blank=True,
    showErrorMessage=True, errorTitle="類型", error="請從下拉選單選擇類型。",
)
status_dv = DataValidation(
    type="list", formula1=f"=選項!$B$2:$B${len(STATUSES) + 1}", allow_blank=True,
    showErrorMessage=True, errorTitle="狀態", error="狀態只能是：現行、已修訂、作廢。",
)
date_dv = DataValidation(
    type="date", operator="greaterThan", formula1="DATE(1900,1,1)", allow_blank=True,
    showErrorMessage=True, errorTitle="日期", error="請輸入日期，例如 2025/3/1。",
)
ws.add_data_validation(type_dv)
ws.add_data_validation(status_dv)
ws.add_data_validation(date_dv)
type_dv.add(f"C2:C{LAST_ROW}")
status_dv.add(f"G2:G{LAST_ROW}")
date_dv.add(f"E2:F{LAST_ROW}")

wb.active = 0
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(OUT)
