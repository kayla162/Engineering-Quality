"""Three exports of the same read-only snapshot, including complete evidence."""
import csv
from io import BytesIO, StringIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate

from app.reports.fonts import ensure_font
from app.schemas.integrated_checklist import IntegratedChecklistOut

SOURCE_LABELS = {
    "both": "人工確認兩邊都有", "needs_confirmation": "標單工項待確認圖說",
    "drawing_only": "人工確認圖說獨有", "drawing_candidate": "圖說候選待確認",
}
BASIS_LABELS = {
    "OK": "已有規範依據（非全面合規）",
    "SPEC_NOT_INCLUDED": "本案規範未收錄", "NEEDS_CONFIRMATION": "請確認依據",
    None: "未判定（圖說候選）",
}
HEADERS = ["識別碼", "標單工項ID", "項次", "名稱", "圖說來源狀態", "規範依據狀態",
           "章碼", "本案章節", "規範依據", "圖說與人工狀態", "工法狀態", "工法候選"]
_XML_ILLEGAL = dict.fromkeys(list(range(9)) + [11, 12] + list(range(14, 32)))
REVIEW_LABELS = {"pending": "待確認", "confirmed": "人工已確認", "rejected": "人工已排除",
                 "stale": "確認已過期"}


def safe_cell(value):
    """Neutralize formulas, including control/whitespace-prefixed payloads."""
    if value is None:
        return ""
    text = str(value)
    if text.startswith(("\t", "\r", "\n")) or text.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _preview(value, compact):
    return value[:180] + "…（完整內容見證據全文）" if compact and len(value) > 180 else value


def _drawing_text(drawing, *, compact=False):
    lines = [
        f"{drawing.filename or '文件名稱未取得'}｜第 {drawing.page} 頁｜"
        f"{drawing.number or '圖號未定'} {drawing.title or '圖名未定'}",
        f"人工狀態：{REVIEW_LABELS.get(drawing.review_status, drawing.review_status)}",
        f"圖項證據：{_preview(drawing.evidence, compact)}",
    ]
    if drawing.review_note:
        lines.append(f"人工備註：{_preview(drawing.review_note, compact)}")
    if drawing.warning:
        lines.append(f"圖頁注意：{_preview(drawing.warning, compact)}")
    if not compact:
        lines.extend([f"圖項ID：{drawing.id}｜文件ID：{drawing.document_id}｜圖頁ID：{drawing.sheet_id}",
                      f"圖框原文：{drawing.sheet_evidence}"])
    return "\n".join(lines)


def _basis_text(basis):
    return f"{basis.reference}｜類型：{basis.basis_type}｜出處：{basis.detail or '未提供'}｜" \
           f"查證：{basis.verification_status or '未查證'}"


def row_values(row, *, compact=False):
    return [
        row.key, str(row.work_item_id or ""), row.item_no or "", row.name,
        SOURCE_LABELS[row.source_status], BASIS_LABELS[row.basis_result],
        row.chapter_code or "",
        f"{row.chapter.code} {row.chapter.title}（版本 {row.chapter.version}）" if row.chapter else "",
        "\n".join(_preview(_basis_text(basis), compact) for basis in row.bases),
        "\n\n".join(_drawing_text(drawing, compact=compact) for drawing in row.drawings),
        "未指定（有待確認候選）" if row.method.candidates else "未指定",
        "\n".join(f"{candidate.text}｜來源：{candidate.source}｜"
                  f"證據：{_preview(candidate.evidence, compact)}" for candidate in row.method.candidates),
    ]


def _metadata(report):
    return [
        ["專案", str(report.project_id)], ["圖說版本", str(report.revision)],
        ["總列數", str(report.total)], ["契約與限制", report.contract_note],
        *[["注意事項", warning] for warning in report.warnings],
        ["摘要", f"標單工項 {report.summary.work_items}；已配對 {report.summary.matched}；"
         f"標單待確認 {report.summary.needs_confirmation}；圖說候選 {report.summary.drawing_candidates}；"
         f"圖說獨有 {report.summary.drawing_only}；已排除 {report.summary.rejected_items}；"
         f"待確認圖項 {report.summary.pending_items}；"
         f"已有依據 {report.summary.basis_ok}；本案未收錄 {report.summary.basis_spec_not_included}；"
         f"依據待確認 {report.summary.basis_needs_confirmation}"],
    ]


def build_csv(report: IntegratedChecklistOut) -> bytes:
    buffer = StringIO(newline="")
    writer = csv.writer(buffer)
    for row in _metadata(report):
        writer.writerow([safe_cell(value) for value in row])
    writer.writerow([])
    writer.writerow(HEADERS)
    for row in report.items:
        writer.writerow([safe_cell(value) for value in row_values(row)])
    return buffer.getvalue().encode("utf-8-sig")


def _leaves(value, path=""):
    """Leaf evidence retains original field names and ordering for reconstruction."""
    if isinstance(value, dict):
        for key, item in value.items():
            yield from _leaves(item, f"{path}.{key}" if path else key)
    elif isinstance(value, list):
        for index, item in enumerate(value):
            yield from _leaves(item, f"{path}.{index}")
    elif isinstance(value, str):
        yield path, value


def _xlsx_text(value):
    # XML 1.0 cannot contain these controls. Encode visibly instead of dropping.
    return "".join(f"\\u{ord(char):04x}" if ord(char) in _XML_ILLEGAL else char for char in str(value))


def build_xlsx(report: IntegratedChecklistOut) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "整合檢核"
    sheet.append(HEADERS)
    evidence = book.create_sheet("證據全文")
    evidence.append(["識別碼", "欄位路徑", "段次", "總段數", "原文"])
    notes = book.create_sheet("說明")
    for metadata in _metadata(report):
        notes.append([safe_cell(_xlsx_text(value)) for value in metadata])
    notes.append(["長文字", "主表超過 Excel 儲存格限制的內容，依識別碼與欄位路徑見證據全文。"])
    notes.append(["安全編碼", "危險公式字首加單引號；XML 不支援的控制字元顯示為 \\uXXXX。"])
    for row in report.items:
        values = []
        for value in row_values(row, compact=True):
            value = safe_cell(_xlsx_text(value))
            if len(value) > 32000:
                value = value[:31000] + "\n[完整原文見「證據全文」，依識別碼查閱]"
            values.append(value)
        sheet.append(values)
        for path, value in _leaves(row.model_dump(mode="json")):
            # Leave room for safety prefix and preserve every character, even
            # when a single field exceeds Excel's 32767-character cell limit.
            value = _xlsx_text(value)
            chunks = [value[start:start + 30000] for start in range(0, len(value), 30000)] or [""]
            for index, chunk in enumerate(chunks, 1):
                evidence.append([row.key, path, index, len(chunks), safe_cell(chunk)])
    for ws in (sheet, evidence):
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="234E52")
        for cells in ws.iter_rows(min_row=2):
            for cell in cells:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for index in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(index)].width = 32 if index > 3 else 24
    evidence.column_dimensions["E"].width = 100
    for index in range(2, sheet.max_row + 1):
        sheet.row_dimensions[index].height = 88
    notes.column_dimensions["A"].width = 20
    notes.column_dimensions["B"].width = 100
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def build_pdf(report: IntegratedChecklistOut) -> bytes:
    font = ensure_font()
    buffer = BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=16 * mm,
                                 rightMargin=16 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
                                 title="工程整合檢核（待人工確認）")
    body = ParagraphStyle("IntegratedBody", fontName=font, fontSize=9, leading=13,
                          wordWrap="CJK", splitLongWords=True, spaceAfter=3)
    heading = ParagraphStyle("IntegratedHeading", parent=body, fontSize=11, leading=16,
                             spaceBefore=8, spaceAfter=4)
    story = [Paragraph("工程整合檢核", heading)]

    def add(label, value, style=body):
        # Small paragraphs split across pages. No unbreakable table rows and
        # no raw user HTML passed to reportlab's mini-XML parser.
        value = _xlsx_text(value)
        chunks = [value[start:start + 2000] for start in range(0, len(value), 2000)] or [""]
        for index, chunk in enumerate(chunks):
            text = f"{label}{'（續）' if index else ''}：{chunk}"
            story.append(Paragraph(escape(text).replace("\n", "<br/>"), style))

    for label, value in _metadata(report):
        add(label, value)
    sheet_appendix = {}
    for index, row in enumerate(report.items, 1):
        add(str(index), f"{row.item_no or '圖說候選'} {row.name}", heading)
        add("識別碼", row.key)
        add("狀態", f"{SOURCE_LABELS[row.source_status]}｜{BASIS_LABELS[row.basis_result]}｜"
            f"工法未指定{'（有待確認候選）' if row.method.candidates else ''}")
        if row.chapter:
            add("本案章節", f"{row.chapter.code} {row.chapter.title}（版本 {row.chapter.version}）")
        elif row.chapter_code:
            add("章碼", row.chapter_code)
        for basis in row.bases:
            add("規範依據", _basis_text(basis))
        for drawing in row.drawings:
            if drawing.sheet_id not in sheet_appendix:
                sheet_appendix[drawing.sheet_id] = (len(sheet_appendix) + 1, drawing)
            appendix_number = sheet_appendix[drawing.sheet_id][0]
            add("圖說來源", f"{drawing.filename or drawing.document_id} 第 {drawing.page} 頁｜"
                f"{drawing.number or '圖號未定'} {drawing.title or '圖名未定'}｜"
                f"{REVIEW_LABELS.get(drawing.review_status, drawing.review_status)}｜圖框全文見附錄 {appendix_number}")
            add("圖項證據", drawing.evidence)
            if drawing.review_note:
                add("人工備註", drawing.review_note)
            if drawing.warning:
                add("圖頁注意", drawing.warning)
            if str(drawing.id) not in row.key:
                add("關聯圖項", str(drawing.id))
        for candidate in row.method.candidates:
            add("待確認工法候選", f"{candidate.text}｜來源：{candidate.source}")
            # Avoid repeating evidence already printed under this exact item.
            if candidate.evidence not in [row.name, *[d.evidence for d in row.drawings],
                                           *[d.title for d in row.drawings]]:
                add("工法候選原文", candidate.evidence)
    if sheet_appendix:
        story.append(PageBreak())
        story.append(Paragraph("圖框證據附錄（同張圖頁只列一次）", heading))
        for number, drawing in sheet_appendix.values():
            add(f"附錄 {number}", f"{drawing.filename or drawing.document_id} 第 {drawing.page} 頁｜"
                f"{drawing.number or '圖號未定'} {drawing.title or '圖名未定'}", heading)
            add("文件識別", f"{drawing.document_id}｜圖頁識別：{drawing.sheet_id}")
            add("圖框原文", drawing.sheet_evidence)

    def page_number(canvas, doc):
        canvas.setFont(font, 8)
        canvas.drawRightString(A4[0] - 16 * mm, 8 * mm, f"第 {doc.page} 頁")

    document.build(story, onFirstPage=page_number, onLaterPages=page_number)
    return buffer.getvalue()
