"""Excel 匯出（openpyxl）。輸入只有 ReportData，不碰資料庫。"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.reports.columns import headers, row_values
from app.schemas.report import ReportData

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
ATTENTION_FILL = PatternFill("solid", fgColor="FFF2CC")


def build(report: ReportData) -> bytes:
    workbook = Workbook()
    _summary_sheet(workbook.active, report)
    _detail_sheet(workbook.create_sheet("法規引用清單"), report)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _summary_sheet(sheet, report: ReportData) -> None:
    sheet.title = "摘要"
    rows = [
        ("文件名稱", report.filename),
        ("文件狀態", report.document_status),
        ("報告產生時間", report.generated_at.strftime("%Y-%m-%d %H:%M:%S %Z")),
        ("法規引用總數", report.total_references),
        ("", ""),
        ("現行有效", report.summary.confirmed_active),
        ("已修正", report.summary.possible_amended),
        ("已廢止", report.summary.possible_repealed),
        ("疑似對應（需人工確認）", report.summary.possible_match),
        ("查無此法規", report.summary.not_found),
        ("查詢失敗或尚未查詢", report.summary.unknown),
    ]
    for label, value in rows:
        sheet.append([label, value])
    for cell in sheet["A"]:
        cell.font = Font(bold=True)

    if report.narrative:
        sheet.append([])
        sheet.append(["摘要說明", report.narrative])
    sheet.append([])
    sheet.append(["資料來源聲明", report.disclaimer])
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 90
    for row in sheet.iter_rows(min_col=2, max_col=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def _detail_sheet(sheet, report: ReportData) -> None:
    sheet.append(headers())
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL

    for row in report.rows:
        sheet.append(row_values(row))
        if row.needs_attention:
            for cell in sheet[sheet.max_row]:
                cell.fill = ATTENTION_FILL

    sheet.freeze_panes = "A2"
    widths = [18, 24, 14, 10, 8, 20, 10, 24, 18, 14, 12, 12, 12, 18, 40, 30]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
