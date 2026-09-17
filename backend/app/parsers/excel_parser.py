"""Excel 解析（.xls 走 xlrd、.xlsx 走 openpyxl）。

為什麼需要這個：工程標單是 Excel，而且真實檔案是 **OLE2 格式的 .xls**——
不是 .xlsx，openpyxl 讀不了。標單是工項清單的來源，沒有這個解析器，
「工項檢核表」整條路就斷在入口（見 docs/superpowers/specs/2026-09-11-工項檢核表-design.md）。

**一列一個 section，不是一格一個。** DOCX 的表格是一格一段，Excel 照抄的話
真實標單（2669 列 x 7 欄）會產生近兩萬個 section，而且把「壹.一.1.1.1」跟
「施工測量，基地測量及鑑界費」拆散。法規抽取跑在完整的列上準得多。

**邊界：這裡不認得標單的結構。** 項次、計價代碼、哪張工作表是詳細表——那些是
工項抽取的職責。它會自己去讀同一份原始檔（上傳的檔案有保留），需要的是有欄位
邊界的結構化讀取，不是這裡產生的文字。把標單格式寫進解析器的話，換一種標單
就得改解析器。
"""

from pathlib import Path

from app.core.enums import ContentType
from app.core.errors import UnprocessableDocumentError
from app.parsers.base import ParsedSection

#: 同一列的儲存格用這個串起來。用非中文字元，抽取時才不會把兩格黏成一個法規名稱。
_CELL_SEPARATOR = " | "


def _format(value: object) -> str:
    """儲存格值轉字串。

    整數型的浮點數要去掉小數尾巴——這不是美觀問題：標準編號、章碼、數量在
    試算表裡常常是數值型別，「3090.0」會讓「CNS 3090」抽不出來。
    """
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _filled(cells: list[object]) -> list[tuple[int, str]]:
    """回傳 [(欄索引, 文字)]，只留有內容的格子。"""
    return [(index, text) for index, text in enumerate(_format(c) for c in cells) if text]


def _is_continuation(current: list[tuple[int, str]], previous: list[tuple[int, str]]) -> bool:
    """這一列是不是前一列的換行續寫？

    標單的長說明會跨列續寫，而且是**從詞中間斷開**的（真實資料）：
        列18  項次=''  「新北市建築物施工中管制要點及新北市建」
        列19  項次=''  「築工程施工圍籬綠美化推動實施計畫規定」

    各自成段的話，「新北市建築工程施工圍籬綠美化推動實施計畫」會被切成兩半，
    抽取器就吐出「市建築物施工中管制要點」「勞安標準」這種半截詞。

    判準刻意不綁標單格式（這裡不該認得「項次」欄是什麼）：
    **只有一欄有內容，而且和前一列是同一欄**——那就是換行造成的，不是新的一筆。
    """
    return (
        len(current) == 1
        and len(previous) >= 1
        and current[0][0] in {column for column, _ in previous}
    )


def _to_section(texts: list[str], order: int, sheet_name: str) -> ParsedSection:
    return ParsedSection(
        content=_CELL_SEPARATOR.join(texts),
        order_index=order,
        content_type=ContentType.TABLE,
        # 試算表沒有頁碼，跟 DOCX 一樣留 None，不要編一個出來
        page=None,
        section_title=sheet_name,
    )


class ExcelParser:
    def parse(self, path: Path) -> list[ParsedSection]:
        rows = (
            self._read_xls(path)
            if path.suffix.lower() == ".xls"
            else self._read_xlsx(path)
        )

        sections: list[ParsedSection] = []
        # 續寫判斷只在同一張工作表內有效，換表就重來
        previous_sheet: str | None = None
        previous_cells: list[tuple[int, str]] = []

        for sheet_name, cells in rows:
            filled = _filled(cells)
            if not filled:
                continue

            same_sheet = sheet_name == previous_sheet
            if sections and same_sheet and _is_continuation(filled, previous_cells):
                # 從詞中間斷開的，中間不能插分隔符號，否則名稱還是斷的
                sections[-1].content += filled[0][1]
            else:
                sections.append(
                    _to_section([text for _, text in filled], len(sections), sheet_name)
                )
                previous_cells = filled
            previous_sheet = sheet_name

        if not sections:
            raise UnprocessableDocumentError("試算表裡沒有任何內容")
        return sections

    # --- 各格式的讀取 ---

    def _read_xlsx(self, path: Path) -> list[tuple[str, list[object]]]:
        from openpyxl import load_workbook

        try:
            # data_only=True 取公式的快取值；沒有快取值時是 None，會被當成空格
            book = load_workbook(path, data_only=True, read_only=True)
        except Exception as exc:
            raise UnprocessableDocumentError(f"無法開啟 Excel 檔案：{exc}") from exc

        try:
            return [
                (sheet.title, list(row))
                for sheet in book.worksheets
                for row in sheet.iter_rows(values_only=True)
            ]
        finally:
            book.close()

    def _read_xls(self, path: Path) -> list[tuple[str, list[object]]]:
        import xlrd

        try:
            book = xlrd.open_workbook(str(path))
        except Exception as exc:
            raise UnprocessableDocumentError(f"無法開啟 Excel 檔案：{exc}") from exc

        return [
            (sheet.name, sheet.row_values(index))
            for sheet in book.sheets()
            for index in range(sheet.nrows)
        ]
