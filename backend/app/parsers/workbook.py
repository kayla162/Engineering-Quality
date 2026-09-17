"""低階活頁簿讀取（.xls 走 xlrd、.xlsx 走 openpyxl）。

ExcelParser 與標單讀取器共用這一層。和 ExcelParser 的差別：這裡**保留欄位位置**
（空格留成 ""），因為標單要知道「第 0 欄是項次、第 6 欄是計價代碼」。
"""

from dataclasses import dataclass
from pathlib import Path

from app.core.errors import UnprocessableDocumentError


@dataclass(frozen=True, slots=True)
class SheetRow:
    sheet: str
    #: 該工作表內的列號（0 起算），換工作表會重新計數
    index: int
    #: 已格式化的字串，**保留欄位位置**，空格是 ""
    cells: list[str]


def format_cell(value: object) -> str:
    """儲存格值轉字串。

    整數型的浮點數要去掉小數尾巴——這不是美觀問題：標準編號、章碼、數量在
    試算表裡常常是數值型別，「3090.0」會讓「CNS 3090」抽不出來。
    """
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_workbook(path: Path) -> list[SheetRow]:
    if path.suffix.lower() == ".xls":
        return _read_xls(path)
    return _read_xlsx(path)


def _read_xlsx(path: Path) -> list[SheetRow]:
    from openpyxl import load_workbook

    try:
        # data_only=True 取公式的快取值；沒有快取值時是 None，會變成空字串
        book = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise UnprocessableDocumentError(f"無法開啟 Excel 檔案：{exc}") from exc

    try:
        return [
            SheetRow(sheet=sheet.title, index=index, cells=[format_cell(c) for c in row])
            for sheet in book.worksheets
            for index, row in enumerate(sheet.iter_rows(values_only=True))
        ]
    finally:
        book.close()


def _read_xls(path: Path) -> list[SheetRow]:
    import xlrd

    try:
        book = xlrd.open_workbook(str(path))
    except Exception as exc:
        raise UnprocessableDocumentError(f"無法開啟 Excel 檔案：{exc}") from exc

    return [
        SheetRow(
            sheet=sheet.name,
            index=index,
            cells=[format_cell(v) for v in sheet.row_values(index)],
        )
        for sheet in book.sheets()
        for index in range(sheet.nrows)
    ]
