import csv
from io import BytesIO, StringIO
import uuid

import openpyxl
import pymupdf
import pytest


def report(name="防水工程", evidence="工法：熱熔工法；待確認"):
    from app.schemas.integrated_checklist import IntegratedChecklistOut
    project, item, sheet, doc = (uuid.uuid4() for _ in range(4))
    return IntegratedChecklistOut.model_validate({
        "project_id": project, "revision": 4, "total": 1,
        "summary": {"total": 1, "drawing_candidates": 1},
        "warnings": ["無標單，不能判定缺項"], "contract_note": "契約優先順序依本案確認",
        "items": [{
            "key": f"drawing:{item}", "name": name, "source_status": "drawing_candidate",
            "drawings": [{"id": item, "sheet_id": sheet, "document_id": doc,
                          "filename": "契約圖.pdf", "page": 7, "number": "A-07",
                          "title": "防水詳圖", "evidence": evidence,
                          "sheet_evidence": "完整圖框原文", "review_status": "pending",
                          "review_note": "人工備註"}],
        }],
    })


@pytest.mark.parametrize("name", ["=1+1", "+SUM(A1)", "-1+1", "@SUM(A1)", "\t=1+1", "\r=1+1"])
def test_csv_and_xlsx_neutralize_formula_injection(name):
    from app.reports.integrated_checklist import build_csv, build_xlsx
    data = report(name=name)
    rows = list(csv.reader(StringIO(build_csv(data).decode("utf-8-sig"))))
    assert any("'" + name in row for row in rows)
    workbook = openpyxl.load_workbook(BytesIO(build_xlsx(data)))
    assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)
    assert any(cell.value == "'" + name for row in workbook["整合檢核"] for cell in row)


def test_exports_retain_full_unicode_evidence_and_pdf_handles_long_escaped_text():
    from app.reports.integrated_checklist import build_csv, build_xlsx, build_pdf
    evidence = "證據 <script> & 原文 " * 2500 + "證據末尾"
    data = report(evidence=evidence)
    assert evidence in build_csv(data).decode("utf-8-sig")
    workbook = openpyxl.load_workbook(BytesIO(build_xlsx(data)))
    # Excel cells are capped at 32767 characters: evidence uses ordered chunks.
    chunks = [row[4] for row in workbook["證據全文"].iter_rows(min_row=2, values_only=True)
              if row[1] == "drawings.0.evidence"]
    assert "".join(chunks) == evidence
    document = pymupdf.open(stream=build_pdf(data), filetype="pdf")
    assert document.page_count > 1
    text = "".join(page.get_text() for page in document)
    assert "證據末尾" in text and "契約圖.pdf" in text
    assert "未指定" in text


def test_method_candidates_require_explicit_text_and_keep_source():
    from app.services.integrated_checklist_service import method_candidates
    assert method_candidates("屋頂防水詳圖", source="drawing:1") == []
    assert method_candidates("不得採用熱熔工法", source="drawing:1") == []
    result = method_candidates("工法：熱熔工法；厚度另核", source="drawing:1")
    assert result[0].text == "熱熔工法"
    assert result[0].source == "drawing:1"
    assert result[0].evidence == "工法：熱熔工法；厚度另核"


@pytest.mark.parametrize("text", [
    "不得採用之工法：熱熔工法",
    "禁止施工工法：熱熔工法",
    "本案不採用施工方法：熱熔工法",
    "禁用工法：熱熔工法",
    "工法：不得採用熱熔工法",
])
def test_negated_method_labels_never_become_selectable_candidates(text):
    from app.services.integrated_checklist_service import method_candidates
    assert method_candidates(text, source="drawing:1") == []


@pytest.mark.parametrize("boundary", ["。", "；", "\n", "!"])
def test_negation_in_previous_sentence_does_not_hide_explicit_method(boundary):
    from app.services.integrated_checklist_service import method_candidates
    result = method_candidates(
        f"不得採用其他工法{boundary}施工工法：熱熔工法", source="drawing:1")
    assert [item.text for item in result] == ["熱熔工法"]


def test_pdf_uses_chinese_labels_and_deduplicates_sheet_evidence():
    from app.reports.integrated_checklist import build_pdf, build_xlsx
    data = report()
    second = data.items[0].model_copy(deep=True)
    second.key = f"drawing:{uuid.uuid4()}"
    second.name = "同頁另一候選"
    data.items.append(second)
    data.total = 2
    data.summary.total = 2
    data.summary.drawing_candidates = 2
    document = pymupdf.open(stream=build_pdf(data), filetype="pdf")
    text = "".join(page.get_text() for page in document)
    assert text.count("完整圖框原文") == 1
    assert "drawings.0." not in text
    assert "人工備註" in text and "圖項證據" in text
    book = openpyxl.load_workbook(BytesIO(build_xlsx(data)))
    summary = book["整合檢核"]["J2"].value
    assert "A-07" in summary and "第 7 頁" in summary
    assert '"sheet_evidence"' not in summary
